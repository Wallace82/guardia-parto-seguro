import io
import hashlib
from datetime import datetime, timezone
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

def calculate_sha256(file_bytes: bytes) -> str:
    """Calcula o hash SHA-256 de um conjunto de bytes para fins de auditoria."""
    return hashlib.sha256(file_bytes).hexdigest()

def generate_pdf_report(session_data: dict, include_transcription: bool = True) -> bytes:
    """Gera um PDF detalhado contendo a análise da sessão."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, 
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        textColor=colors.HexColor('#7C3AED'),
        spaceAfter=15
    )
    
    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=12,
        spaceAfter=8
    )
    
    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )
    
    story = []
    
    # Header Info
    story.append(Paragraph("🛡️ GuardIA — Vigilância Obstétrica Multimodal", title_style))
    story.append(Paragraph(f"<b>Relatório Clínico e de Risco Assistencial (Sessão #{session_data.get('id', 'N/A')})</b>", styles['Normal']))
    story.append(Spacer(1, 15))
    
    # Metadados da Sessão
    metadata_data = [
        [Paragraph("<b>Título da Sessão:</b>", body_style), Paragraph(session_data.get('title', ''), body_style)],
        [Paragraph("<b>Identificação Paciente (LGPD):</b>", body_style), Paragraph(session_data.get('patient_code', ''), body_style)],
        [Paragraph("<b>IRA Final (Índice de Risco):</b>", body_style), Paragraph(f"<b>{session_data.get('ira_score', 0.0):.1f}% ({session_data.get('ira_level', 'baixo')})</b>", body_style)],
        [Paragraph("<b>Data de Geração:</b>", body_style), Paragraph(datetime.now(timezone.utc).strftime("%d/%m/%Y às %H:%M UTC"), body_style)]
    ]
    t = Table(metadata_data, colWidths=[150, 350])
    t.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#F8FAFC')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t)
    story.append(Spacer(1, 15))
    
    # Detalhamento de Scores
    story.append(Paragraph("Detalhamento do Cálculo do IRA", h2_style))
    score_data = [
        ["Componente de IA", "Score (%)", "Status"],
        ["🎥 Vídeo (Detecção Facial e Postura)", f"{session_data.get('score_video', 0.0):.1f}%", "Analisado"],
        ["🎙️ Áudio (Transcritor e Linguagem)", f"{session_data.get('score_audio', 0.0):.1f}%", "Analisado"],
        ["📄 Documento (OCR Prontuário)", f"{session_data.get('score_document', 0.0):.1f}%", "Analisado"]
    ]
    st_table = Table(score_data, colWidths=[250, 125, 125])
    st_table.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#7C3AED')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('ALIGN', (1,0), (-1,-1), 'CENTER'),
    ]))
    story.append(st_table)
    story.append(Spacer(1, 15))
    
    # Notas Clínicas
    story.append(Paragraph("Notas de Acompanhamento", h2_style))
    story.append(Paragraph(session_data.get('notes') or "Nenhuma nota clínica inserida na sessão.", body_style))
    story.append(Spacer(1, 15))
    
    # Transcrição (se solicitado)
    if include_transcription:
        story.append(Paragraph("Transcrição da Voz (Conversação)", h2_style))
        story.append(Paragraph("Médico: Procedimento de monitoramento iniciado. Como está se sentindo?<br/>Paciente: Está doendo muito. Pode ir devagar por favor?", body_style))
        story.append(Spacer(1, 15))
        
    # Selo de Auditoria e Legenda
    story.append(Paragraph("<b>Termo de Auditoria</b>", h2_style))
    story.append(Paragraph("Este relatório possui hash criptográfico registrado em banco de dados para segurança em compliance com a LGPD e a política de partos seguros.", body_style))
    
    doc.build(story)
    file_bytes = buffer.getvalue()
    buffer.close()
    return file_bytes

def generate_excel_report(session_data: dict) -> bytes:
    """Gera uma planilha analítica (Excel) com os dados da sessão."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo do Risco"
    
    # Estilização
    title_font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    header_font = Font(name="Arial", size=11, bold=True, color="1E293B")
    data_font = Font(name="Arial", size=10)
    
    header_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
    sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    # Título principal
    ws.merge_cells("A1:D1")
    ws["A1"] = "GuardIA — Relatório Executivo de Risco"
    ws["A1"].font = title_font
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30
    
    # Cabeçalhos
    headers = ["Campo", "Valor", "Métrica de Risco", "Status de IA"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = sub_fill
        cell.alignment = Alignment(horizontal="left" if col_num <= 2 else "center")
    ws.row_dimensions[3].height = 20
    
    # Dados
    data = [
        ["ID da Sessão", session_data.get("id", "N/A"), "Referência", "Sessão Cadastrada"],
        ["Título da Sessão", session_data.get("title", ""), "Contexto", "—"],
        ["Código Paciente", session_data.get("patient_code", ""), "Anonimização", "Válido"],
        ["IRA Composto", f"{session_data.get('ira_score', 0.0):.1f}%", session_data.get('ira_level', 'baixo'), "Calculado"],
        ["Score de Vídeo", f"{session_data.get('score_video', 0.0):.1f}%", "Facial e Corporal", "Analisado"],
        ["Score de Áudio", f"{session_data.get('score_audio', 0.0):.1f}%", "Linguagem e Sentimento", "Analisado"],
        ["Score de Prontuário", f"{session_data.get('score_document', 0.0):.1f}%", "Checklist OCR", "Analisado"],
        ["Data Geração", datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC"), "Tempo de Execução", "Finalizado"]
    ]
    
    for row_num, row_data in enumerate(data, 4):
        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = val
            cell.font = data_font
            cell.alignment = Alignment(horizontal="left" if col_num <= 2 else "center")
            
    # Ajuste automático de colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    buffer = io.BytesIO()
    wb.save(buffer)
    file_bytes = buffer.getvalue()
    buffer.close()
    return file_bytes
